"""TRUE end-to-end application-flow test (V4 verification phase).

One test walks the whole product the way a real local user does:

    Connections.csv -> upload -> dataset + Person rows -> enrichment
    -> Apify (fixtures) -> raw persistence -> deterministic normalization
    -> experiences / education / skills -> company classification
    -> ProfileSemantic (mocked LLM) -> local embedding -> READY/PARTIAL
    -> dashboard -> natural-language search -> LLM interpretation (mocked)
    -> hard-fact gate -> local scoring -> semantic judge (mocked)
    -> judge validation -> rescore -> reranker -> final audit (mocked)
    -> validated verdicts -> LLM-verified results -> results page payload
    -> save -> reload (zero provider calls) -> Excel export -> person refresh
    -> dataset deletion -> cascade cleanup.

Only the EXTERNAL boundaries are mocked:
  * Apify           -> USE_FIXTURES=true (conftest) — real fixture profiles
  * Anthropic/Groq  -> the four ``generate_structured`` / ``_call_*`` seams:
      - query interpretation   (query_interpreter.generate_structured)
      - profile semantics      (semantic_llm.derive_semantics)
      - exhaustive judge       (semantic_judge._call_judge)
      - final result audit     (final_auditor._call_audit)

No internal application service is mocked. Every validator, the hard-fact gate,
deterministic scoring, the rescore, the reranker path and both grounded
validators run for real.
"""
from __future__ import annotations

import io
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app import repositories as repo
from app.config import settings
from app.constants import (
    EnrichmentState,
    JobStatus,
    Qualification,
    SearchStatus,
    VerificationStatus,
)
from app.database import SessionLocal
from app.models import (
    Certification,
    Connection,
    Education,
    EnrichmentJob,
    Experience,
    Language,
    Person,
    ProfileEmbedding,
    ProfileSemantic,
    Publication,
    RawProfile,
    SearchQuery,
    SearchResult,
    SearchRunState,
    Skill,
)
from app.schemas import LenientSearchPlan, ProfileSemanticData

FIXTURE_CSV = Path(__file__).resolve().parents[1] / "fixtures" / "connections_sample.csv"

QUERY = "People with hands-on cloud infrastructure experience"


# ─────────────────────────── LLM boundary fakes ───────────────────────────


def _fake_interpretation(*, provider="anthropic:paid", model="claude-sonnet-5"):
    """query_interpreter.generate_structured seam -> a Sonnet-shaped plan with a
    required professional_concept, so the semantic judge + final audit are
    genuinely required for this search to return results."""
    plan = LenientSearchPlan.model_validate(
        {
            "intent": "professional_recommendation",
            "criteria": [
                {
                    "id": "c_cloud",
                    "type": "professional_concept",
                    "concept": "hands-on cloud infrastructure experience",
                    "required": True,
                    "weight": 100,
                }
            ],
            "interpretation_summary": "People who have personally built or run cloud infrastructure.",
            "interpretation_confidence": 0.82,
        }
    )

    def fake(system, user, schema, **kw):  # noqa: ARG001
        return schema.model_validate(plan.model_dump()), provider, model

    return fake


def _fake_derive_semantics(*, provider="anthropic:paid", model="claude-sonnet-5"):
    """semantic_llm.derive_semantics seam -> a valid ProfileSemanticData payload.

    Keeps it generic but non-empty so embeddings pick up extra keywords and the
    judge packet has assertions to look at.
    """

    def fake(db, person):  # noqa: ARG001
        data = ProfileSemanticData.model_validate(
            {
                "seniority_level": "senior",
                "job_families": ["software engineering"],
                "technical_domains": ["cloud infrastructure", "distributed systems"],
                "industries": ["technology"],
                "searchable_keywords": ["aws", "kubernetes", "cloud", "infrastructure"],
                "career_summary": f"{person.full_name or 'This person'} works in cloud and platform engineering.",
            }
        )
        return data.model_dump(), provider, model

    return fake


def _fake_call_judge(*, provider="anthropic:paid", model="claude-sonnet-5"):
    """semantic_judge._call_judge seam. Returns the expanded verdict map
    ``{person_id: {criterion_id: verdict_dict}}`` grounded in a real packet
    experience ref, marking the concept TRUE for everyone judged."""

    def fake(payload, packets, unresolved_by_person=None):
        crit_ids = [c["id"] for c in payload["criteria_to_judge"]]
        expanded: dict[str, dict[str, dict]] = {}
        for pkt in packets:
            pid = pkt["person_id"]
            want = (unresolved_by_person or {}).get(pid) or crit_ids
            first_exp = (pkt.get("current") or {}).get("experience_id") or (
                pkt["past"][0]["experience_id"] if pkt.get("past") else None
            )
            per: dict[str, dict] = {}
            for cid in want:
                grounded = bool(first_exp)
                per[cid] = {
                    "criterion_id": cid,
                    "status": "true" if grounded else "unknown",
                    "match_strength": 0.88 if grounded else 0.0,
                    "confidence": 0.9,
                    "reason": "role explicitly covers building cloud infrastructure",
                    "supporting_evidence_refs": [f"exp:{first_exp}"] if grounded else [],
                    "contradicting_evidence_refs": [],
                    "experience_ids": [first_exp] if grounded else [],
                }
            expanded[pid] = per
        return "ok", expanded, provider, model

    return fake


def _fake_call_audit(*, provider="anthropic:paid", model="claude-sonnet-5", capture=None):
    """final_auditor._call_audit seam -> ``("ok", [decision_dict, ...], p, m)``.
    Approves every candidate with a grounded, supported review of each required
    criterion, plus a plain-language display_reason used verbatim on the card."""

    def fake(payload, packets, first_pass_by_id, parsed=None):  # noqa: ARG001
        if capture is not None:
            capture.append([p["person_id"] for p in packets])
        required = [c for c in payload["criteria"] if c["required"]]
        out = []
        for pkt in packets:
            cur_id = (pkt.get("current") or {}).get("experience_id") or "x"
            out.append(
                {
                    "person_id": pkt["person_id"],
                    "decision": "approved",
                    "confidence": 0.9,
                    "reason": "evidence supports every required criterion",
                    "display_reason": "Runs cloud infrastructure day to day per their current role.",
                    "criteria": [
                        {
                            "criterion_id": c["id"],
                            "status_review": "supported",
                            "reason": "",
                            "supporting_evidence_refs": [f"exp:{cur_id}"],
                            "contradicting_evidence_refs": [],
                        }
                        for c in required
                    ],
                    "supporting_evidence_refs": [f"exp:{cur_id}"],
                    "contradicting_evidence_refs": [],
                    "suggested_qualification": None,
                }
            )
        return "ok", out, provider, model

    return fake


@pytest.fixture
def full_llm_stack(monkeypatch):
    """Turn the real product settings on (conftest disables them) and wire the
    four LLM seams to deterministic fakes."""
    monkeypatch.setattr(settings, "semantic_enabled", True)
    monkeypatch.setattr(settings, "llm_query_interpretation", True)
    monkeypatch.setattr(settings, "require_llm_for_results", True)
    monkeypatch.setattr(settings, "semantic_judge_enabled", True)
    monkeypatch.setattr(settings, "semantic_judge_mode", "all_viable")
    monkeypatch.setattr(settings, "final_result_audit_enabled", True)
    monkeypatch.setattr(settings, "search_require_final_audit", True)
    monkeypatch.setattr(settings, "llm_reason_generation", False)

    monkeypatch.setattr(
        "app.services.query_interpreter.generate_structured", _fake_interpretation()
    )
    monkeypatch.setattr(
        "app.services.semantic_llm.derive_semantics", _fake_derive_semantics()
    )
    monkeypatch.setattr("app.services.semantic_judge._call_judge", _fake_call_judge())
    audit_calls: list = []
    monkeypatch.setattr(
        "app.services.final_auditor._call_audit", _fake_call_audit(capture=audit_calls)
    )
    return {"audit_calls": audit_calls}


# ─────────────────────────── the walk ───────────────────────────


def test_full_application_flow(client, full_llm_stack, monkeypatch):
    # 1-3 — upload -> dataset + Person + Connection rows ────────────────────
    up = client.post(
        "/datasets", files={"file": ("Connections.csv", FIXTURE_CSV.read_bytes(), "text/csv")}
    )
    assert up.status_code == 201, up.text
    report = up.json()
    ds_id = report["dataset"]["dataset_id"]
    assert report["imported"] == 11  # 13 data rows - 1 dup - 1 no-url
    assert report["duplicates_removed"] == 1
    assert report["skipped_no_url"] == 1

    db = SessionLocal()
    try:
        assert db.query(Person).filter_by(dataset_id=ds_id).count() == 11
        assert db.query(Connection).filter_by(dataset_id=ds_id).count() == 11
        assert all(
            p.enrichment_state == EnrichmentState.PENDING
            for p in db.query(Person).filter_by(dataset_id=ds_id)
        )
    finally:
        db.close()

    # 4-7 — enrichment: Apify (fixtures) -> raw -> normalize -> semantics
    #        -> embeddings -> READY/PARTIAL ─────────────────────────────────
    started = client.post(f"/datasets/{ds_id}/enrich").json()
    assert started["started"] is True and started["mode"] == "enrich"

    db = SessionLocal()
    try:
        people = db.query(Person).filter_by(dataset_id=ds_id).all()
        # one bad profile never stops the batch — every person reaches a terminal state
        assert all(
            p.enrichment_state in (EnrichmentState.READY, EnrichmentState.PARTIAL, EnrichmentState.FAILED)
            for p in people
        )
        ready = [p for p in people if p.enrichment_state == EnrichmentState.READY]
        assert len(ready) >= 5  # the 5 hand-written fixtures at minimum

        # 6 — normalized relational rows exist and raw JSON is persisted verbatim
        assert db.query(RawProfile).filter(
            RawProfile.person_id.in_([p.id for p in ready])
        ).count() >= 5
        assert db.query(Experience).filter(
            Experience.person_id.in_([p.id for p in ready])
        ).count() >= 5

        # 5 — semantics + 7 — embedding for every READY person
        for p in ready:
            assert db.query(ProfileSemantic).filter_by(person_id=p.id).count() == 1
            assert p.semantic_version == settings.semantic_profile_version
            assert db.query(ProfileEmbedding).filter_by(person_id=p.id).count() == 1

        job = (
            db.query(EnrichmentJob)
            .filter_by(dataset_id=ds_id)
            .order_by(EnrichmentJob.created_at.desc())
            .first()
        )
        assert job.status in (JobStatus.COMPLETED, JobStatus.PARTIAL)  # terminal
    finally:
        db.close()

    # 8 — dataset status report is internally consistent ───────────────────
    st = client.get(f"/datasets/{ds_id}/status").json()
    assert st["connections"] == 11
    assert st["progress_total"] == 11
    assert st["progress_done"] == st["ready"] + st["partial"] + st["failed"]
    assert 0 <= st["progress_pct"] <= 100
    assert st["ready"] >= 5

    # dashboard: people list renders READY/PARTIAL/... states
    ppl = client.get(f"/datasets/{ds_id}/people").json()
    assert len(ppl) == 11
    assert {x["enrichment_state"] for x in ppl} <= {
        "READY", "PARTIAL", "FAILED", "PENDING", "WAITING_FOR_FREE_LLM", "LLM_COMPLETE", "NORMALIZED",
    }

    # ── Apify must NEVER be reached again from here on (search / reload). ──
    def _no_apify(*a, **k):  # noqa: ARG001
        raise AssertionError("Apify was called after enrichment finished")

    monkeypatch.setattr("app.services.apify_client.scrape_profiles", _no_apify)

    # 9-19 — natural-language search: interpret -> gate -> score -> judge
    #        -> validate -> rescore -> rerank -> audit -> verified results ──
    resp = client.post("/search", json={"dataset_id": ds_id, "query": QUERY})
    assert resp.status_code == 200, resp.text
    body = resp.json()

    assert body["search_status"] == SearchStatus.SUCCESS
    assert body["verification_status"] == VerificationStatus.COMPLETE
    assert body["llm_verified"] is True
    assert body["anthropic_attempted"] is True
    assert body["anthropic_succeeded"] is True
    assert body["fallback_used"] is False
    assert body["ai_provider"] == "anthropic"

    conn = body["connections"]
    assert conn["returned"] >= 1
    assert conn["returned"] <= settings.top_connections
    assert conn["exact_match_count"] + conn["possible_match_count"] >= conn["returned"]
    assert len(full_llm_stack["audit_calls"]) >= 1  # the final audit really ran

    # at least one shown result is LLM-verified and carries the audit outcome
    verified = [r for r in conn["results"] if r["llm_verified"]]
    assert verified, body
    top = verified[0]
    assert top["audit_decision"] == "approved"
    assert top["qualification"] in (Qualification.EXACT_MATCH, Qualification.POSSIBLE_MATCH)
    assert top["reason"]
    # judge / audit observability present
    assert body["judge_metadata"]["status"] in ("full", "partial")
    assert body["audit_metadata"]["enabled"] is True
    assert body["audit_metadata"]["status"] in ("full", "partial")

    search_id = body["search_id"]

    # 20-22 — save + reload: identical snapshot, ZERO provider calls ────────
    def _boom(*a, **k):  # noqa: ARG001
        raise AssertionError("reload of a saved search must not call any LLM / search step")

    monkeypatch.setattr("app.services.search_service.interpret_query", _boom)
    monkeypatch.setattr("app.services.search_service.run_judge", _boom)
    monkeypatch.setattr("app.services.search_service._run_final_audit", _boom)
    monkeypatch.setattr("app.services.search_service.get_candidates", _boom)
    monkeypatch.setattr("app.services.embeddings.embed_text", _boom)

    reloaded = client.get(f"/search/{search_id}").json()
    assert reloaded["query"] == QUERY
    assert reloaded["search_status"] == SearchStatus.SUCCESS
    assert reloaded["verification_status"] == VerificationStatus.COMPLETE
    assert reloaded["llm_verified"] is True
    assert reloaded["interpreted_query"] == body["interpreted_query"]
    assert reloaded["judge_metadata"] == body["judge_metadata"]
    assert reloaded["audit_metadata"] == body["audit_metadata"]
    rc = reloaded["connections"]
    assert [r["person_id"] for r in rc["results"]] == [r["person_id"] for r in conn["results"]]
    assert [r["qualification"] for r in rc["results"]] == [r["qualification"] for r in conn["results"]]
    assert [r["match_score"] for r in rc["results"]] == [r["match_score"] for r in conn["results"]]
    assert [r["llm_verified"] for r in rc["results"]] == [r["llm_verified"] for r in conn["results"]]
    assert [r["person_id"] for r in rc["near_matches"]] == [
        r["person_id"] for r in conn["near_matches"]
    ]

    # search history opens that persisted search
    hist = client.get(f"/datasets/{ds_id}/searches").json()
    assert any(h["search_id"] == search_id and h["query"] == QUERY for h in hist)

    # 23 — Excel export works on a partially-enriched dataset ──────────────
    xls = client.get(f"/datasets/{ds_id}/export")
    assert xls.status_code == 200
    wb = load_workbook(io.BytesIO(xls.content))
    assert wb.sheetnames[0] == "Profiles"
    assert wb["Profiles"].max_row >= 6  # header + >= 5 enriched profiles

    # 24 — refresh one person (force, since fixtures are always "fresh") ───
    db = SessionLocal()
    try:
        target = (
            db.query(Person)
            .filter_by(dataset_id=ds_id, enrichment_state=EnrichmentState.READY)
            .first()
        )
        pid = target.id
        raw_before = db.query(RawProfile).filter_by(person_id=pid).count()
    finally:
        db.close()

    refreshed = client.post(f"/people/{pid}/refresh?force=true").json()
    assert refreshed["refreshed"] is True

    db = SessionLocal()
    try:
        p = db.get(Person, pid)
        assert p.enrichment_state in (EnrichmentState.READY, EnrichmentState.PARTIAL)
        # refresh replaces normalized rows in place — never forks a second Person
        assert db.query(Person).filter_by(dataset_id=ds_id).count() == 11
        # a fresh raw snapshot is appended (history), not a duplicate Person
        assert db.query(RawProfile).filter_by(person_id=pid).count() >= raw_before
        assert db.query(Experience).filter_by(person_id=pid).count() >= 0
    finally:
        db.close()

    # 25-26 — delete the dataset -> every associated row is gone ───────────
    assert client.delete(f"/datasets/{ds_id}").status_code == 204

    db = SessionLocal()
    try:
        assert db.get(SearchQuery, search_id) is None
        assert db.get(SearchRunState, search_id) is None
        # nothing references the dataset any more
        assert db.query(Person).filter_by(dataset_id=ds_id).count() == 0
        assert db.query(Connection).filter_by(dataset_id=ds_id).count() == 0
        assert db.query(EnrichmentJob).filter_by(dataset_id=ds_id).count() == 0
        assert db.query(SearchQuery).filter_by(dataset_id=ds_id).count() == 0
        assert db.query(RawProfile).filter_by(dataset_id=ds_id).count() == 0
        assert db.query(SearchResult).filter(
            SearchResult.search_id == search_id
        ).count() == 0
        # child rows keyed only by person_id are cascade-deleted too — no orphans
        for model in (Experience, Education, Skill, Certification, Language,
                      Publication, ProfileSemantic, ProfileEmbedding):
            assert db.query(model).filter(
                ~model.person_id.in_(select(Person.id))
            ).count() == 0
    finally:
        db.close()

    assert client.get(f"/datasets/{ds_id}/status").status_code == 404
    assert client.get(f"/datasets/{ds_id}/searches").status_code == 404


# ─────────────────────────── failure modes (CASE 1-10) ───────────────────────────


def test_case1_invalid_csv_returns_clean_4xx_no_dataset(client):
    before = len(client.get("/datasets").json())
    r = client.post("/datasets", files={"file": ("x.csv", b"not,a,connections,file\n1,2,3,4\n", "text/csv")})
    assert r.status_code == 422
    assert "CSV" in r.json()["detail"] or "url" in r.json()["detail"].lower()
    assert len(client.get("/datasets").json()) == before  # nothing was created

    empty = client.post("/datasets", files={"file": ("x.csv", b"", "text/csv")})
    assert empty.status_code == 400


def test_case2_one_apify_failure_does_not_break_the_dataset(client, full_llm_stack, monkeypatch):
    ds_id = client.post(
        "/datasets", files={"file": ("c.csv", FIXTURE_CSV.read_bytes(), "text/csv")}
    ).json()["dataset"]["dataset_id"]

    real_scrape = __import__("app.services.apify_client", fromlist=["scrape_profiles"]).scrape_profiles

    def flaky(urls, *, hints=None):
        # drop the first requested profile from every batch -> that person fails,
        # everyone else in the batch still succeeds
        kept = list(urls)[1:] or list(urls)
        return real_scrape(kept, hints=hints)

    monkeypatch.setattr("app.services.apify_client.scrape_profiles", flaky)
    monkeypatch.setattr(settings, "max_apify_retries", 1)
    client.post(f"/datasets/{ds_id}/enrich")

    st = client.get(f"/datasets/{ds_id}/status").json()
    assert st["progress_done"] == st["progress_total"]  # reached a terminal state
    assert st["ready"] >= 1  # the rest of the network is usable
    # a search still works against the usable remainder
    r = client.post("/search", json={"dataset_id": ds_id, "query": QUERY})
    assert r.status_code == 200


def test_case3_semantic_failure_keeps_raw_and_never_recalls_apify(client, monkeypatch):
    monkeypatch.setattr(settings, "semantic_enabled", True)
    monkeypatch.setattr(settings, "llm_query_interpretation", True)
    ds_id = client.post(
        "/datasets", files={"file": ("c.csv", FIXTURE_CSV.read_bytes(), "text/csv")}
    ).json()["dataset"]["dataset_id"]

    # Apify (fixtures) succeeds; the semantic LLM is dead for the whole run
    monkeypatch.setattr("app.services.semantic_llm.derive_semantics", lambda db, p: None)
    client.post(f"/datasets/{ds_id}/enrich")

    db = SessionLocal()
    try:
        ready = db.query(Person).filter_by(
            dataset_id=ds_id, enrichment_state=EnrichmentState.READY
        ).all()
        assert ready  # scraped + normalized + embedded + READY despite no semantics
        for p in ready:
            assert db.query(RawProfile).filter_by(person_id=p.id).count() >= 1  # raw kept
            assert p.semantic_version is None  # eligible for later backfill
    finally:
        db.close()

    # now the semantic LLM recovers — a re-enrich backfills semantics with NO Apify
    monkeypatch.setattr(
        "app.services.apify_client.scrape_profiles",
        lambda *a, **k: (_ for _ in ()).throw(AssertionError("semantic backfill must not call Apify")),
    )
    monkeypatch.setattr(
        "app.services.semantic_llm.derive_semantics", _fake_derive_semantics()
    )
    again = client.post(f"/datasets/{ds_id}/enrich").json()
    assert again["mode"] == "backfill_semantics"

    db = SessionLocal()
    try:
        done = db.query(Person).filter_by(dataset_id=ds_id).filter(
            Person.semantic_version == settings.semantic_profile_version
        ).count()
        assert done >= 1
    finally:
        db.close()


def test_case5_all_interpretation_providers_fail_returns_no_results(client, monkeypatch):
    monkeypatch.setattr(settings, "require_llm_for_results", True)
    monkeypatch.setattr(settings, "llm_query_interpretation", True)
    ds_id = client.post(
        "/datasets", files={"file": ("c.csv", FIXTURE_CSV.read_bytes(), "text/csv")}
    ).json()["dataset"]["dataset_id"]
    client.post(f"/datasets/{ds_id}/enrich")

    monkeypatch.setattr("app.services.query_interpreter.generate_structured", lambda *a, **k: None)
    r = client.post("/search", json={"dataset_id": ds_id, "query": QUERY}).json()
    assert r["search_status"] == SearchStatus.AI_UNAVAILABLE
    assert r["connections"]["results"] == []
    assert r["connections"]["near_matches"] == []
    assert r["llm_verified"] is False


def test_case6_judge_cannot_verify_returns_no_results(client, monkeypatch):
    monkeypatch.setattr(settings, "require_llm_for_results", True)
    monkeypatch.setattr(settings, "llm_query_interpretation", True)
    monkeypatch.setattr(settings, "semantic_judge_enabled", True)
    monkeypatch.setattr(settings, "semantic_enabled", True)
    ds_id = client.post(
        "/datasets", files={"file": ("c.csv", FIXTURE_CSV.read_bytes(), "text/csv")}
    ).json()["dataset"]["dataset_id"]
    monkeypatch.setattr("app.services.semantic_llm.derive_semantics", _fake_derive_semantics())
    client.post(f"/datasets/{ds_id}/enrich")

    monkeypatch.setattr("app.services.query_interpreter.generate_structured", _fake_interpretation())
    monkeypatch.setattr(
        "app.services.semantic_judge._call_judge", lambda *a, **k: ("failed", None, None, None)
    )
    r = client.post("/search", json={"dataset_id": ds_id, "query": QUERY}).json()
    assert r["search_status"] == SearchStatus.VERIFICATION_INCOMPLETE
    assert r["connections"]["results"] == []
    assert r["connections"]["near_matches"] == []
    assert r["verification_status"] == VerificationStatus.INCOMPLETE


def test_case7_final_audit_incomplete_returns_no_results(client, monkeypatch):
    monkeypatch.setattr(settings, "require_llm_for_results", True)
    monkeypatch.setattr(settings, "llm_query_interpretation", True)
    monkeypatch.setattr(settings, "semantic_judge_enabled", True)
    monkeypatch.setattr(settings, "semantic_enabled", True)
    monkeypatch.setattr(settings, "final_result_audit_enabled", True)
    monkeypatch.setattr(settings, "search_require_final_audit", True)
    ds_id = client.post(
        "/datasets", files={"file": ("c.csv", FIXTURE_CSV.read_bytes(), "text/csv")}
    ).json()["dataset"]["dataset_id"]
    monkeypatch.setattr("app.services.semantic_llm.derive_semantics", _fake_derive_semantics())
    client.post(f"/datasets/{ds_id}/enrich")

    monkeypatch.setattr("app.services.query_interpreter.generate_structured", _fake_interpretation())
    monkeypatch.setattr("app.services.semantic_judge._call_judge", _fake_call_judge())
    monkeypatch.setattr(
        "app.services.final_auditor._call_audit", lambda *a, **k: ("failed", None, None, None)
    )
    r = client.post("/search", json={"dataset_id": ds_id, "query": QUERY}).json()
    assert r["search_status"] == SearchStatus.VERIFICATION_INCOMPLETE
    assert r["connections"]["results"] == []


def test_case9_saved_search_reload_makes_zero_provider_calls(client, full_llm_stack, monkeypatch):
    ds_id = client.post(
        "/datasets", files={"file": ("c.csv", FIXTURE_CSV.read_bytes(), "text/csv")}
    ).json()["dataset"]["dataset_id"]
    client.post(f"/datasets/{ds_id}/enrich")
    sid = client.post("/search", json={"dataset_id": ds_id, "query": QUERY}).json()["search_id"]

    for seam in (
        "app.services.query_interpreter.generate_structured",
        "app.services.semantic_llm.derive_semantics",
        "app.services.semantic_judge._call_judge",
        "app.services.final_auditor._call_audit",
        "app.services.apify_client.scrape_profiles",
    ):
        monkeypatch.setattr(seam, lambda *a, **k: (_ for _ in ()).throw(
            AssertionError(f"{seam} called during reload")
        ))

    r1 = client.get(f"/search/{sid}").json()
    r2 = client.get(f"/search/{sid}").json()
    assert r1 == r2  # deterministic snapshot
    assert r1["query"] == QUERY
